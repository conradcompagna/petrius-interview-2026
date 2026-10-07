"""Build the editable two-page Task I note from docs/task1_note.md."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
source = (ROOT / 'docs/task1_note.md').read_text(encoding='utf-8')
doc = Document()
for style in doc.styles:
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.62)
sec.left_margin = sec.right_margin = Inches(0.7)
sec.header_distance = sec.footer_distance = Inches(0.3)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Subtitle'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for node in list(style.element.rPr) if style.element.rPr is not None else []:
        if node.tag == qn('w:color'):
            for att in ('themeColor','themeTint','themeShade'):
                node.attrib.pop(qn('w:'+att), None)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.02
for name, size in [('Title',18),('Heading 1',12),('Heading 2',11)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = name != 'Title'
    s.paragraph_format.space_before = Pt(8 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.keep_with_next = True

lines = source.splitlines(); i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line == '<!-- PAGEBREAK -->':
        doc.add_page_break(); i += 1; continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-:]+',c) for c in cells): rows.append(cells)
            i += 1
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths=[1.95,.85,1.25,1.25,.85,.85]
        for row_no, row in enumerate(rows):
            for col_no, value in enumerate(row):
                cell = table.cell(row_no,col_no)
                cell.width = Inches(widths[col_no])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = value
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_no == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.line_spacing=1
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.bold = row_no == 0
                tcPr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                tcPr.append(borders)
                if row_no==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EDF2');tcPr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
                tcPr.append(margins)
        continue
    else:
        p=doc.add_paragraph(line)
        if line.startswith('Conrad Compagna'):
            p.paragraph_format.space_after=Pt(8)
            p.runs[0].font.size=Pt(10)
        if line.startswith('Images:'):
            for run in p.runs:run.font.size=Pt(9)
    i += 1
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Conrad Compagna  |  Task I  |  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.author='Conrad Compagna'
doc.core_properties.title='Task I page layout method and development'
doc.save(ROOT / 'docs/task1_note.docx')
print('Created task1_note.docx; source words:', len(source.split()))
